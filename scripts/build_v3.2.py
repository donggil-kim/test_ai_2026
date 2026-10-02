"""통합 AI 보안위협 매트릭스 v3.2 LITE 워크북 빌드.

입력: data/v3.2/base.yaml(Lv3) · lv2.yaml · domains.yaml · cases.yaml · changelog.yaml
      data/v3.2/rewrite/*.yaml(v3.2 문체 재작성본, Lv3 ID 키)
      reference/atlas/ATLAS-2026_09.yaml(기법 Maturity)
출력: output/통합_AI보안위협_매트릭스_v3.2_LITE.xlsx

발생가능성·근거 수준·실제 사고 수는 사례 연결과 OWASP 인용 근거에서 계산해 값으로 기록하고,
위험도·개요 지표·도메인 집계는 워크북 수식으로 연결한다.
"""
import re
from pathlib import Path

import openpyxl
import yaml
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data/v3.2"
OUT = ROOT / "output/통합_AI보안위협_매트릭스_v3.2_LITE.xlsx"


def load(name):
    return yaml.safe_load((DATA / name).read_text(encoding="utf-8"))


base, lv2, domains, cases, log = (load(f) for f in
                                  ("base.yaml", "lv2.yaml", "domains.yaml", "cases.yaml", "changelog.yaml"))
rewrite = {}
for f in sorted((DATA / "rewrite").glob("*.yaml")):
    rewrite.update(yaml.safe_load(f.read_text(encoding="utf-8")) or {})
maturity = {k: v.get("maturity") for k, v in
            yaml.safe_load((ROOT / "reference/atlas/ATLAS-2026_09.yaml").read_text(encoding="utf-8"))["techniques"].items()}
lv2_by = {x["ut"]: x for x in lv2}
dom_by = {d["name"][:3]: d for d in domains}
DATE = log["date"]

# ── 행 데이터 조립 ───────────────────────────────────────────────
rows = []
for b in base:
    rw = rewrite.get(b["id"], {})
    r = {**b, **rw}
    if rw and "style" not in rw:
        r["style"] = "v3.2"
    ids = [c for c in cases if r["id"] in c["lv3"]]
    am, an = len(ids), sum(c["type"] == "Incident" for c in ids)
    ab = sum(maturity.get(t) == "Realized" for t in r["atlas"] + r["atlas_related"])
    ao, ap = r["owasp_incidents"], r["owasp_vulns"]
    r["likelihood"] = "상" if an + ao >= 2 else ("중" if (an + ao == 1 or am or ab or ap) else "하")
    r["evidence"] = ("실제 사고 확인" if an + ao else "실사용 기법 포함" if ab else
                     "실증·공개 취약점" if (am or ap) else "이론·시나리오")
    r["incidents"] = an + ao
    r["domain"] = lv2_by[r["ut"]]["domain"]
    r["lv2"] = lv2_by[r["ut"]]["name"]
    r["control"] = lv2_by[r["ut"]]["control"]
    r["owasp_main"] = ", ".join(c for k in ("LLM", "ASI", "DSGAI", "MCP") for c in r["owasp"][k]) or None
    r["atlas_txt"] = ", ".join(t.replace("AML.", "") for t in r["atlas"]) or None
    rows.append(r)
N = len(rows)
LAST = 4 + N
written = sum(r["style"] == "v3.2" for r in rows)

# ── 공통 서식 ────────────────────────────────────────────────────
NAVY = "1F3864"
TITLE = Font(bold=True, size=14, color=NAVY)
SUB = Font(size=9, color="595959")
SECTION = Font(bold=True, size=11, color=NAVY)
TH_FONT, TH_FILL = Font(bold=True, size=9, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
GROUP_FILL = PatternFill("solid", fgColor="D9E1F2")
DOMAIN_FILL = PatternFill("solid", fgColor="FDE9D9")
BODY, BODY_OLD, BODY_B = Font(size=9), Font(size=9, color="8C8C8C"), Font(size=9, bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOP_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
RISK = ('=IF(AND({g}="상",{h}="상"),"매우 높음",IF(OR(AND({g}="상",{h}="중"),AND({g}="중",{h}="상")),"높음",'
        'IF(OR(AND({g}="중",{h}="중"),AND({g}="상",{h}="하"),AND({g}="하",{h}="상")),"보통","낮음")))')

out = openpyxl.Workbook()
out.remove(out.active)


def th(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row, col + i, v)
        c.font, c.fill, c.alignment, c.border = TH_FONT, TH_FILL, CENTER, BOX


def td(ws, row, values, col=1, center=()):
    for i, v in enumerate(values):
        c = ws.cell(row, col + i, v)
        c.font, c.border = BODY, BOX
        c.alignment = CENTER if i in center else WRAP


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths):
        ws.column_dimensions[L(i + 1)].width = w


def lines(text, width):
    """셀 줄 수 추정(한글 1자 = 2폭)."""
    n = 0
    for ln in str(text or "").split("\n"):
        w = sum(2 if ord(ch) > 0x2E80 else 1 for ch in ln)
        n += max(1, -(-w // max(1, int(width * 1.15))))
    return n


# ── 1. 통합매트릭스 ─────────────────────────────────────────────
COLS = [  # (그룹, 열 이름, 너비, 키, 가운데 정렬)
    ("분류 체계 (Lv1 · Lv2 · Lv3)", "도메인 (Lv1)", 13, "domain", False),
    (None, "UT-ID", 7, "ut", True),
    (None, "위협 분류 (Lv2)", 15, "lv2", False),
    (None, "Lv3 ID", 8, "id", True),
    (None, "세부 위협 (Lv3)", 15, "name", False),
    (None, "요약설명", 52, "summary", False),
    (None, "참조", 88, "reference", False),
    ("위험 평가", "발생가능성", 7, "likelihood", True),
    (None, "심각도", 7, "severity", True),
    (None, "위험도", 8, None, True),
    (None, "심각도 근거", 22, "severity_reason", False),
    (None, "근거 수준", 11, "evidence", True),
    (None, "실제 사고 수", 7, "incidents", True),
    ("교차 매핑", "주 공격 단계", 13, "stage", False),
    (None, "ATLAS 기법 (주)", 20, "atlas_txt", False),
    (None, "OWASP 매핑 (주)", 13, "owasp_main", False),
    ("통제 (Lv2)", "핵심 통제 요약", 30, "control", False),
]
col_of = {name: L(j) for j, (_, name, *_) in enumerate(COLS, start=1)}
ws = out.create_sheet("통합매트릭스")
ws["A1"] = f"통합 AI 보안위협 매트릭스 v3.2 LITE — Lv1 도메인(10) · Lv2 위협 분류(36) · Lv3 세부 위협({N})"
ws["A1"].font = TITLE
ws["A2"] = ("기준: MITRE ATLAS v2026.09 · OWASP LLM Top 10 2026 · Agentic Top 10 2026 · GenAI Data Security 2026 v1.0 · "
            f"MCP Top 10(2025 beta) | {DATE} | {log['status']}" + (" — 회색 글자 = v3.1 문구 유지 행" if written < N else ""))
ws["A2"].font = SUB
group_start = None
for j, (group, name, width, *_ ) in enumerate(COLS, start=1):
    if group:
        if group_start and j - 1 > group_start:  # 단일 열 그룹은 병합하지 않음(Excel 복구 경고 방지)
            ws.merge_cells(start_row=3, start_column=group_start, end_row=3, end_column=j - 1)
        group_start = j
        ws.cell(3, j, group)
    c = ws.cell(3, j)
    c.font, c.fill, c.alignment, c.border = Font(bold=True, size=9, color=NAVY), GROUP_FILL, CENTER, BOX
    th(ws, 4, [name], col=j)
    ws.column_dimensions[L(j)].width = width
if len(COLS) > group_start:
    ws.merge_cells(start_row=3, start_column=group_start, end_row=3, end_column=len(COLS))

G, H = col_of["발생가능성"], col_of["심각도"]
for i, r in enumerate(rows):
    row = 5 + i
    for j, (_, name, width, key, center) in enumerate(COLS, start=1):
        v = RISK.format(g=f"{G}{row}", h=f"{H}{row}") if name == "위험도" else r.get(key)
        c = ws.cell(row, j, v)
        c.border = BOX
        c.alignment = TOP_CENTER if center else WRAP
        c.font = BODY_B if name in ("세부 위협 (Lv3)", "위험도") else BODY
        if name in ("요약설명", "참조") and r["style"] != "v3.2":
            c.font = BODY_OLD
        if name == "도메인 (Lv1)":
            c.fill = DOMAIN_FILL
    ws.row_dimensions[row].height = min(409, 12.5 * max(lines(r["summary"], 52), lines(r["reference"], 88), 3) + 4)
ws.freeze_panes = "F5"
ws.auto_filter.ref = f"A4:{L(len(COLS))}{LAST}"
for col in (G, H):
    for k, color in (("상", "C00000"), ("중", "7F6000"), ("하", "375623")):
        ws.conditional_formatting.add(f"{col}5:{col}{LAST}",
                                      CellIsRule(operator="equal", formula=[f'"{k}"'], font=Font(bold=True, color=color)))
R = col_of["위험도"]
for k, bg, fg in (("매우 높음", "C00000", "FFFFFF"), ("높음", "ED7D31", "FFFFFF"),
                  ("보통", "FFE699", "000000"), ("낮음", "C6E0B4", "000000")):
    ws.conditional_formatting.add(f"{R}5:{R}{LAST}", CellIsRule(
        operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", fgColor=bg), font=Font(bold=True, color=fg)))
E = col_of["근거 수준"]
for k, color in (("실제 사고 확인", "C00000"), ("실사용 기법 포함", "C55A11"),
                 ("실증·공개 취약점", "7F6000"), ("이론·시나리오", "595959")):
    ws.conditional_formatting.add(f"{E}5:{E}{LAST}",
                                  CellIsRule(operator="equal", formula=[f'"{k}"'], font=Font(bold=True, color=color)))
dv = DataValidation(type="list", formula1='"상,중,하"', allow_blank=False)
dv.add(f"{G}5:{H}{LAST}")
ws.add_data_validation(dv)

M = "통합매트릭스!"
rng = lambda name: f"{M}${col_of[name]}$5:${col_of[name]}${LAST}"

# ── 2. 개요 ─────────────────────────────────────────────────────
o = out.create_sheet("개요", 0)
o["A1"], o["A1"].font = "통합 AI 보안위협 매트릭스 (UTM) v3.2 LITE", TITLE
o["A2"], o["A2"].font = (f"Lv1 도메인 10 · Lv2 위협 분류 36 · Lv3 세부 위협 {N} · 위험 평가(발생가능성×심각도) | {DATE} | "
                         f"{log['status']}"), SUB
row = 4
o.cell(row, 1, "1. 통합 대상 소스").font = SECTION
row += 1
th(o, row, ["소스", "버전 / 공개일", "규모", "관점", "입수 경로"])
row += 1
for r in [
    ("MITRE ATLAS", "v2026.09 (2026-09-15)", "전술 16 · 기법 208 · 완화책 40 · 사례 73 · 절차 665", "공격자 TTP", "ATLAS-2026_09.yaml · excel-atlas*.xlsx"),
    ("OWASP Top 10 for LLM Applications", "2026 (2026-08-04)", "10", "LLM 애플리케이션 위험", "원문 PDF"),
    ("OWASP Top 10 for Agentic Applications", "2026 (2025-12-09)", "10 + 사고 추적표 25건", "에이전트 위험", "원문 PDF"),
    ("OWASP GenAI Data Security", "2026 v1.0 (2026-03)", "21", "데이터 생애주기 위험", "원문 PDF"),
    ("OWASP MCP Top 10", "2025 (beta)", "10", "MCP 프로토콜 계층", "github.com/OWASP/www-project-mcp-top-10"),
]:
    td(o, row, r)
    row += 1
row += 1
o.cell(row, 1, "2. 핵심 지표 (수식 연동)").font = SECTION
row += 1
th(o, row, ["지표", "값"])
row += 1
for k, v in [
    ("Lv1 도메인 / Lv2 위협 분류 / Lv3 세부 위협",
     f'=SUMPRODUCT(1/COUNTIF({rng("도메인 (Lv1)")},{rng("도메인 (Lv1)")}))&" / "&'
     f'SUMPRODUCT(1/COUNTIF({rng("UT-ID")},{rng("UT-ID")}))&" / "&COUNTA({rng("Lv3 ID")})'),
    ("위험도 매우 높음 / 높음 / 보통 / 낮음",
     "=" + '&" / "&'.join(f'COUNTIF({rng("위험도")},"{k}")' for k in ("매우 높음", "높음", "보통", "낮음"))),
    ("근거 수준 실제 사고 확인 / 실사용 기법 포함 / 실증·공개 취약점 / 이론·시나리오",
     "=" + '&" / "&'.join(f'COUNTIF({rng("근거 수준")},"{k}")'
                          for k in ("실제 사고 확인", "실사용 기법 포함", "실증·공개 취약점", "이론·시나리오"))),
    ("실제 사고가 1건 이상 연결된 Lv3", f'=COUNTIF({rng("실제 사고 수")},">0")'),
    ("v3.2 문체 재작성 Lv3 / 전체", f"{written} / {N}"),
]:
    td(o, row, [k, v])
    row += 1
row += 1
o.cell(row, 1, "3. 위험 매트릭스 (Lv3 수)").font = SECTION
row += 1
th(o, row, ["발생가능성 \\ 심각도", "상", "중", "하"])
row += 1
for g in ("상", "중", "하"):
    td(o, row, [g] + [f'=COUNTIFS({rng("발생가능성")},"{g}",{rng("심각도")},"{h}")' for h in ("상", "중", "하")],
       center=(0, 1, 2, 3))
    row += 1
row += 1
o.cell(row, 1, "4. 시트 구성").font = SECTION
row += 1
th(o, row, ["시트", "내용"])
row += 1
for k, v in [
    ("통합매트릭스", f"핵심 산출물. Lv3 {N}행 — 분류 체계, 요약설명·참조(AI 관점·실제 사례), 위험 평가, 교차 매핑(공격 단계·ATLAS·OWASP), 핵심 통제(Lv2)"),
    ("도메인·평가기준", "3단 분류 정의와 작성 규칙, Lv1 도메인 리스크 정의·경계 판정 규칙, 위험 평가·근거 수준 기준, 참조 작성 규칙"),
    ("ATLAS사례연구", "MITRE ATLAS 사례 73건 요약과 Lv3 매핑 — 참조의 'ATLAS AML.CS' 출처 조회용"),
    ("변경이력", "버전별 변경 내역과 Lv3 ID 대응표(폐지 ID는 재사용하지 않음)"),
]:
    td(o, row, [k, v])
    row += 1
widths(o, [44, 62, 34, 22, 34])

# ── 3. 도메인·평가기준 ──────────────────────────────────────────
d = out.create_sheet("도메인·평가기준")
d["A1"], d["A1"].font = "분류 체계 정의와 위험 평가 기준", TITLE
row = 3
d.cell(row, 1, "1. 3단 분류 체계").font = SECTION
row += 1
th(d, row, ["단계", "정의", "작성 규칙", "개수"])
row += 1
for r in [
    ("Lv1 도메인", "영역별 보안 리스크(공격 흐름상 위협 행위자 영역을 D01로 배치)",
     "OWASP 위협·리스크 명명 규칙 준용: [영역] + [위협 명사], 대등한 두 위협만 '및'으로 결합 / 경계 판정 규칙은 아래 2번 표",
     f'=SUMPRODUCT(1/COUNTIF({rng("도메인 (Lv1)")},{rng("도메인 (Lv1)")}))'),
    ("Lv2 위협 분류", "통합 위협(UT). ID 불변", "'·' 나열 없이 하나의 개념으로 명명",
     f'=SUMPRODUCT(1/COUNTIF({rng("UT-ID")},{rng("UT-ID")}))'),
    ("Lv3 세부 위협", "Lv2 안에서 공격 경로 또는 발생 원리가 독립된 위협",
     "① 독립 메커니즘 ② 근거 1개 이상(ATLAS 기법·OWASP 예시·사례) ③ 공격 단계·생애주기·통제 중 하나 이상 구별 / "
     "병합·폐지된 ID는 재사용하지 않고 신설은 해당 Lv2의 다음 번호 부여", f'=COUNTA({rng("Lv3 ID")})'),
]:
    td(d, row, list(r), center=(3,))
    row += 1
row += 1
d.cell(row, 1, "2. Lv1 도메인 정의").font = SECTION
row += 1
th(d, row, ["도메인", "영문명", "리스크 정의", "주요 대상 영역", "경계 판정 규칙", "Lv2 수", "Lv3 수",
            "위험도 '매우 높음'·'높음' Lv3"])
row += 1
A, B, I_ = rng("도메인 (Lv1)"), rng("UT-ID"), rng("위험도")
for dm in domains:
    a = f"A{row}"
    td(d, row, [dm["name"], dm["en"], dm["definition"], dm["targets"], dm["boundary"],
                f"=SUMPRODUCT(({A}={a})/COUNTIFS({A},{A},{B},{B}))", f"=COUNTIF({A},{a})",
                f'=COUNTIFS({A},{a},{I_},"매우 높음")+COUNTIFS({A},{a},{I_},"높음")'], center=(5, 6, 7))
    row += 1
row += 1
d.cell(row, 1, "3. 위험 평가 기준 (Lv3 단위)").font = SECTION
row += 1
th(d, row, ["항목", "상", "중", "하", "산정 방식"])
row += 1
for r in [
    ("발생가능성", "실제 사고 2건 이상 (ATLAS Incident 사례 + OWASP 인용 실제 사고)",
     "실제 사고 1건, 또는 실증 사례(연구·레드팀)·Realized 기법·공개 취약점/익스플로잇 존재",
     "이론적으로 가능하나 실제 사례·실증·공개 취약점 없음",
     "근거(ATLAS 사례 연결·OWASP 인용 근거·Realized 기법)에서 빌드 시 산정해 값으로 기록. 사례·근거 개정 시 재산정"),
    ("심각도", "AI 시스템 전체 통제권 탈취 또는 중요·개인정보 대량 탈취 가능성 높음 (호스트·에이전트 장악, 연동 자격증명 탈취, 대량 반출)",
     "부분적 통제 상실 또는 제한된 범위의 정보 노출·무결성 훼손·재무 피해",
     "영향 범위 제한적, 경미한 정보 노출 또는 일시적 장애 (준비·탐색 단계 포함)",
     "분석자 평가(심각도 근거 열에 사유 기재) — 기준: 통제권 범위, 정보 노출 규모, 피해 지속성. 드롭다운으로 수정 가능"),
    ("위험도", "매우 높음: 상×상", "높음: 상×중, 중×상 / 보통: 중×중, 상×하, 하×상", "낮음: 중×하, 하×중, 하×하",
     "발생가능성 × 심각도 3×3 매트릭스(수식, 자동 재계산)"),
]:
    td(d, row, list(r))
    row += 1
row += 1
d.cell(row, 1, "4. 근거 수준 정의 (높은 순)").font = SECTION
row += 1
th(d, row, ["근거 수준", "판정 기준"])
row += 1
for r in [
    ("실제 사고 확인", "ATLAS Incident 사례 또는 OWASP 원문이 인용한 실제 사고가 1건 이상 연결('실제 사고 수' ≥ 1)"),
    ("실사용 기법 포함", "실제 사고 연결은 없으나 ATLAS Maturity 'Realized' 기법이 귀속(주·연관 기법)"),
    ("실증·공개 취약점", "연구·레드팀 실증 사례(ATLAS Exercise) 또는 OWASP 인용 취약점·익스플로잇 존재 — CVE, 벤더가 확인·조치한 취약점, 운영 서비스·모델 대상 공개 PoC 포함"),
    ("이론·시나리오", "공개 근거 없음(OWASP 위험 정의·시나리오 기반). '발생하지 않음'이 아니라 '공개 근거 없음'을 의미"),
]:
    td(d, row, list(r))
    row += 1
row += 1
d.cell(row, 1, "5. 요약설명·참조 작성 규칙").font = SECTION
row += 1
th(d, row, ["항목", "규칙"])
row += 1
for r in [
    ("요약설명", "2줄 개조식. 1줄: '공격자는 [수단·경로]로 [행위]할 수 있음' / 2줄: 피해·확산 특성 또는 탐지·방어가 어려운 이유"),
    ("참조 — ■ AI 관점", "3~4줄. 대상·경로, 수법 변형, 연계 Lv3(UT-ID), 범위 경계를 기술하고 출처를 괄호로 표기 — 예: (ATLAS AML.T0068), (OWASP LLM01:2026 예시 #5)"),
    ("참조 — ■ 실제 사례", "1~3줄. [실제 사고] > [공개 취약점] > [실증] > [시나리오] 순으로 선정, '사례명(YYYY-MM): 경위·결과 (출처)' 형식. 근거가 없으면 '공개 사고 미확인 — 사유'"),
    ("사례 라벨", "[실제 사고] 실제 발생 사고 · [공개 취약점] CVE·벤더 확인 취약점·운영 서비스 대상 PoC · [실증] 연구·레드팀 시연(ATLAS Exercise 등) · [시나리오] OWASP 공격 시나리오"),
]:
    td(d, row, list(r))
    row += 1
widths(d, [24, 30, 48, 40, 48, 8, 8, 14])

# ── 4. ATLAS사례연구 ────────────────────────────────────────────
a = out.create_sheet("ATLAS사례연구")
a["A1"], a["A1"].font = "MITRE ATLAS 사례연구(73건) — 요약 및 Lv3 매핑", TITLE
a["A2"], a["A2"].font = ("주 Lv3 = 분석자가 지정한 대표 세부 위협. 매핑 Lv3 = 공식 절차(employs) 기법의 주 Lv3 + 분석자 보완. "
                         "유형: Incident=실제 사고, Exercise=실증"), SUB
th(a, 4, ["사례 ID", "사례명", "유형", "발생·공개 시점", "행위자", "대상", "국문 요약", "주 Lv3", "매핑 Lv3 (전체)", "URL"])
for i, c in enumerate(cases):
    td(a, 5 + i, [c["id"], c["name"], c["type"], c["date"], c["actor"], c["target"], c["summary"], c["main_lv3"],
                  ", ".join(c["lv3"]), c["url"]], center=(2, 3, 7))
a.freeze_panes = "B5"
a.auto_filter.ref = f"A4:J{4 + len(cases)}"
widths(a, [12, 34, 9, 10, 18, 22, 48, 8, 28, 30])

# ── 5. 변경이력 ─────────────────────────────────────────────────
h = out.create_sheet("변경이력")
h["A1"], h["A1"].font = "변경이력", TITLE
h["A2"], h["A2"].font = "개정은 새 버전 파일로 저장하고 이 시트 상단에 기록. v3.1 이전 상세 이력은 v3.1 원본 '변경이력' 시트 참조", SUB
row = 4
h.cell(row, 1, f"1. v3.2 변경 내역 ({DATE})").font = SECTION
row += 1
th(h, row, ["버전", "구분", "대상", "변경 내용", "사유"])
row += 1
for e in log["entries"]:
    td(h, row, e, center=(0, 1))
    row += 1
row += 1
h.cell(row, 1, "2. Lv3 ID 대응 (v3.1 → v3.2)").font = SECTION
row += 1
th(h, row, ["v3.1 ID", "v3.1 명칭", "v3.2 ID", "v3.2 명칭", "변경 유형", "사유"])
row += 1
for e in log["lv3_map"]:
    td(h, row, e, center=(0, 2, 4))
    row += 1
row += 1
h.cell(row, 1, "3. 이전 버전").font = SECTION
row += 1
th(h, row, ["버전", "일자", "구분", "변경 내용", "사유"])
row += 1
for e in log["previous"]:
    td(h, row, e, center=(0, 1, 2))
    row += 1
widths(h, [12, 24, 14, 70, 48, 48])

for s in out.worksheets:
    s.sheet_view.showGridLines = s.title == "통합매트릭스"
out.save(OUT)
print(OUT.relative_to(ROOT), f"Lv3 {N} (v3.2 문체 {written})")
