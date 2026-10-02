"""v3.1 → v3.2 분류 체계 이관 (1회성).

v3.1 원본 워크북을 읽어 v3.2 분류 체계 변경(병합·폐지·신설·명칭 변경·기법 재귀속)을
적용하고 data/v3.2/base.yaml(Lv3 행), lv2.yaml, domains.yaml, cases.yaml을 생성한다.
신규 문체 재작성본은 data/v3.2/rewrite/*.yaml에 별도로 두고 빌드 시 병합한다.
"""
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "reference/utm/통합_AI보안위협_매트릭스_v3.1.xlsx"
ATLAS = ROOT / "reference/atlas/ATLAS-2026_09.yaml"
OUT = ROOT / "data/v3.2"

# ── v3.2 분류 체계 변경 ──────────────────────────────────────────
# 폐지 Lv3 → 흡수 Lv3 (폐지 ID는 재사용하지 않음)
RETIRE = {
    "UT-03.4": "UT-02.1",  # 공격용 에이전트 도구 확보 → 공격 도구 확보
    "UT-07.3": "UT-10.1",  # 검색 노출용 콘텐츠 제작 → 검색 코퍼스 오염
    "UT-08.3": "UT-08.2",  # 런타임 권한 탐색 → 에이전트 구성 및 권한 탐색
    "UT-09.2": "UT-28.1",  # 오염 데이터셋 유포 → 변조 모델 및 데이터셋 유입
    "UT-20.2": "UT-28.3",  # 환각 악용 공격 → 환각 패키지 선점
}
# 기법 주 Lv3 재귀속 (폐지 Lv3 소속 기법은 RETIRE로 자동 이동)
TECH_MOVE = {
    "AML.T0065": "UT-06.1",      # LLM Prompt Crafting: 난독화 → 시스템 지시 무력화
    "AML.T0063": "UT-01.4",      # Discover AI Model Outputs: 부채널 추론 → 모델 및 환경 탐색
    "AML.T0052.001": "UT-04.3",  # Deepfake-Assisted Phishing: AI 생성 피싱 → 신뢰 주체 사칭
}
# 신설 Lv3 (해당 Lv2의 다음 미사용 번호)
NEW = {
    "UT-15.4": dict(after="UT-15.3", name="학습 데이터 추출", derived_from="UT-12.1",
                    stage="유출(Exfiltration)", severity="중",
                    severity_reason="암기된 학습 데이터 단편 노출(대량 추출은 반복 질의 필요)",
                    owasp=dict(LLM=["LLM02"], DSGAI=["DSGAI18"], DSGAI_sub=["DSGAI01"]),
                    owasp_incidents=0, owasp_vulns=1),
    "UT-22.4": dict(after="UT-22.3", name="도구 인자 명령 주입", derived_from="UT-22.1",
                    stage="실행(Execution)", severity="상",
                    severity_reason="도구 실행 호스트의 임의 명령 실행(시스템 장악)",
                    owasp=dict(ASI=["ASI05"], ASI_sub=["ASI02"], LLM_sub=["LLM01", "LLM10"],
                               MCP=["MCP05"]),
                    owasp_incidents=0, owasp_vulns=3),
}
# 신설 Lv3의 연관 기법(주 귀속은 기존 Lv3 유지, Realized 근거 산정에 포함)
NEW_RELATED = {"UT-22.4": ["AML.T0050"], "UT-15.4": ["AML.T0057"]}
# 사례 → Lv3 분석자 보완 추가 / 주 Lv3 변경
SUPP_ADD = {"AML.CS0062": ["UT-22.4"], "AML.CS0052": ["UT-22.4"]}
CASE_MAIN = {"AML.CS0062": "UT-22.4"}
# 명칭 변경
RENAME_LV2 = {
    "UT-02": "공격 준비 및 일반 침투",
    "UT-15": "모델 기반 데이터 추론",
    "UT-28": "AI 개발 공급망 침해",
    "UT-32": "단말 AI 어시스턴트 악용",
}
RENAME_LV3 = {
    "UT-04.1": "딥페이크 생체인증 우회",
    "UT-06.3": "프롬프트 난독화",
    "UT-08.2": "에이전트 구성 및 권한 탐색",
    "UT-12.1": "응답 경유 민감정보 유출",
    "UT-25.1": "평문 저장 자격증명 수집",
    "UT-25.3": "인증 토큰 및 세션 쿠키 탈취",
    "UT-26.2": "캐시된 자격증명 재사용",
    "UT-28.1": "변조 모델 및 데이터셋 유입",
}
# Lv2 핵심 통제 요약 보완(범위가 바뀐 Lv2)
CONTROL = {
    "UT-15": "차등프라이버시 학습·학습 데이터 중복 제거, 출력 확률·점수 비공개, 쿼리량 제한, 추출·멤버십 추론 감사",
    "UT-22": "코드 생성과 실행 분리·승인 게이트, 도구 인자 검증·파라미터화(셸·eval 직접 전달 금지), 비루트 격리 샌드박스, 실행 전 정적검사, 명령 허용목록",
    "UT-28": "AIBOM/SBOM 인벤토리, 버전 고정·출처 검증, AI 권고 패키지 실존·평판 확인, 공급자 약관·라이선스 검토, 레지스트리 허용목록",
}
# 신설·병합으로 OWASP 인용 근거를 재배분하는 Lv3 (사고 수, 취약점·익스플로잇 수)
OWASP_COUNT = {
    "UT-22.1": (0, 3),  # Framelink Figma MCP RCE·VS Code 에이전트 RCE → UT-22.4로 이동
    "UT-15.3": (0, 1),  # Whisper Leak(암호화 스트리밍 부채널, 공개 익스플로잇) — T0063 이동으로 ATLAS 실증 사례 제외
    "UT-06.3": (0, 1),  # M365 Copilot ASCII 스머글링 PoC(비가시 유니코드 = 텍스트 채널 은닉)를 UT-06.4에서 이동
    "UT-06.4": (0, 1),  # 근거 교체: 의료 영상 멀티모달 인젝션(OWASP LLM01:2026 시나리오 #6)
}

OWASP_KEYS = ["LLM", "LLM_sub", "ASI", "ASI_sub", "DSGAI", "DSGAI_sub", "MCP", "MCP_sub"]


def split(v):
    return [x.strip() for x in str(v).replace("·", ",").split(",") if x and x.strip()] if v else []


def block_str(dumper, data):
    style = "|" if "\n" in data else None
    return dumper.represent_scalar("tag:yaml.org,2002:str", data, style=style)


yaml.add_representer(str, block_str)


def dump(obj, path):
    path.write_text(yaml.dump(obj, allow_unicode=True, sort_keys=False, width=4096), encoding="utf-8")


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    atlas = yaml.safe_load(ATLAS.read_text(encoding="utf-8"))
    tech_main = {r[0]: r[6] for r in wb["ATLAS기법매핑"].iter_rows(min_row=5, values_only=True) if r[0]}
    for t, lv3 in list(tech_main.items()):
        if lv3 in RETIRE:
            tech_main[t] = RETIRE[lv3]
    tech_main.update(TECH_MOVE)

    rows = [r for r in wb["통합매트릭스"].iter_rows(min_row=5, values_only=True) if r[3]]
    lv2, base, by_id = {}, [], {}
    for r in rows:
        ut = r[1]
        lv2.setdefault(ut, dict(ut=ut, domain=r[0], name=RENAME_LV2.get(ut, r[2]),
                                prev_name=r[2] if ut in RENAME_LV2 else None,
                                control=CONTROL.get(ut, r[46])))
        owasp = {k: split(r[29 + i]) for i, k in enumerate(OWASP_KEYS)}
        row = dict(id=r[3], ut=ut, name=RENAME_LV3.get(r[3], r[4]),
                   prev=dict(id=r[3], name=r[4]),
                   summary=r[5], reference=r[44], style="v3.1",
                   severity=r[7], severity_reason=r[10], stage=r[11],
                   atlas=[], atlas_related=["AML." + t for t in split(r[25])], owasp=owasp,
                   owasp_incidents=int(r[40] or 0), owasp_vulns=int(r[41] or 0))
        by_id[r[3]] = row
        base.append(row)

    # 폐지 Lv3를 흡수 Lv3에 병합(OWASP 매핑 합집합, OWASP 인용 근거 합산)
    for old, new in RETIRE.items():
        src, dst = by_id[old], by_id[new]
        for k in OWASP_KEYS:
            dst["owasp"][k] = sorted(set(dst["owasp"][k]) | set(src["owasp"][k]))
        dst["atlas_related"] = sorted(set(dst["atlas_related"]) | set(src["atlas_related"]))
        dst["owasp_incidents"] += src["owasp_incidents"]
        dst["owasp_vulns"] += src["owasp_vulns"]
        dst.setdefault("merged", []).append(dict(id=old, name=src["name"]))
        base.remove(src)
    for row in base:  # 주 매핑에 있는 코드는 보조에서 제거
        for k in ("LLM", "ASI", "DSGAI", "MCP"):
            row["owasp"][k + "_sub"] = [c for c in row["owasp"][k + "_sub"] if c not in row["owasp"][k]]

    for lv3, (inc, vul) in OWASP_COUNT.items():
        by_id[lv3]["owasp_incidents"], by_id[lv3]["owasp_vulns"] = inc, vul

    for nid, spec in NEW.items():
        owasp = {k: spec["owasp"].get(k, []) for k in OWASP_KEYS}
        row = dict(id=nid, ut=nid[:5], name=spec["name"], prev=None, derived_from=spec["derived_from"],
                   summary=None, reference=None, style="v3.2",
                   severity=spec["severity"], severity_reason=spec["severity_reason"], stage=spec["stage"],
                   atlas=[], atlas_related=NEW_RELATED.get(nid, []), owasp=owasp,
                   owasp_incidents=spec["owasp_incidents"], owasp_vulns=spec["owasp_vulns"])
        base.insert(base.index(by_id[spec["after"]]) + 1, row)
        by_id[nid] = row

    for t, lv3 in tech_main.items():  # 기법 주 Lv3 (영향 패싯 제외)
        if lv3 in by_id:
            by_id[lv3]["atlas"].append(t)
    for row in base:
        row["atlas_related"] = [t for t in row["atlas_related"] if t not in row["atlas"]]
        row["atlas"].sort(key=lambda t: [int(x) if x.isdigit() else x for x in t.replace("AML.T", "").split(".")])

    # 사례 → Lv3: 공식 절차 기법의 주 Lv3 + 분석자 보완(폐지 ID는 흡수 Lv3로 치환)
    rel = atlas["relationships"]
    v31_tech = {r[0]: r[6] for r in wb["ATLAS기법매핑"].iter_rows(min_row=5, values_only=True) if r[0]}
    cases = []
    for r in wb["ATLAS사례연구"].iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        steps = rel.get(r[0], {}).get("employs", [])
        derived_v31 = {v31_tech[e["target"]] for e in steps if str(v31_tech[e["target"]]).startswith("UT-")}
        supp = {x.strip() for x in (r[9] or "").split(",") if x.strip()} - derived_v31
        supp = {RETIRE.get(x, x) for x in supp} | set(SUPP_ADD.get(r[0], []))
        derived = {tech_main[e["target"]] for e in steps if str(tech_main[e["target"]]).startswith("UT-")}
        lv3s = sorted(derived | supp, key=lambda x: [int(p) for p in x[3:].split(".")])
        main_lv3 = CASE_MAIN.get(r[0], RETIRE.get(r[7], r[7]))
        cases.append(dict(id=r[0], name=r[1], type=r[2], date=str(r[3]), actor=r[4], target=r[5],
                          summary=r[6], main_lv3=main_lv3, lv3=lv3s, supplements=sorted(supp), url=r[15]))

    domains = []
    for r in wb["도메인·평가기준"].iter_rows(min_row=11, max_row=20, max_col=5, values_only=True):
        domains.append(dict(name=r[0], en=r[1], definition=r[2], targets=r[3], boundary=r[4]))

    OUT.mkdir(parents=True, exist_ok=True)
    dump(base, OUT / "base.yaml")
    dump(list(lv2.values()), OUT / "lv2.yaml")
    dump(domains, OUT / "domains.yaml")
    dump(cases, OUT / "cases.yaml")
    print(f"Lv3 {len(base)} · Lv2 {len(lv2)} · 사례 {len(cases)}")


if __name__ == "__main__":
    main()
