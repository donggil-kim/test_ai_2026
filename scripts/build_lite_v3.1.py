"""통합 AI 보안위협 매트릭스 v3.1 원본에서 LITE 최종본을 생성한다.

- 원본(reference/utm/...v3.1.xlsx)은 수정하지 않는다.
- 삭제된 시트를 참조하던 수식은 원본에 캐시된 계산값으로 고정한다.
- 위험도·개요 지표·도메인 집계는 LITE 내부 시트만 참조하는 수식으로 다시 연결한다.
"""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "reference/utm/통합_AI보안위협_매트릭스_v3.1.xlsx"
OUT = ROOT / "output/통합_AI보안위협_매트릭스_v3.1_LITE.xlsx"
DATE = "2026-10-02"

fx = openpyxl.load_workbook(SRC)                   # 서식·수식
val = openpyxl.load_workbook(SRC, data_only=True)  # 캐시 값
out = openpyxl.Workbook()
out.remove(out.active)

NAVY = "1F3864"
TITLE = Font(bold=True, size=14, color=NAVY)
SUB = Font(size=9, color="595959")
SECTION = Font(bold=True, size=11, color=NAVY)
TH_FONT = Font(bold=True, size=9, color="FFFFFF")
TH_FILL = PatternFill("solid", fgColor=NAVY)
GROUP_FILL = PatternFill("solid", fgColor="D9E1F2")
BODY = Font(size=9)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

RISK = '=IF(AND({g}="상",{h}="상"),"매우 높음",IF(OR(AND({g}="상",{h}="중"),AND({g}="중",{h}="상")),"높음",' \
       'IF(OR(AND({g}="중",{h}="중"),AND({g}="상",{h}="하"),AND({g}="하",{h}="상")),"보통","낮음")))'


def th(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row, col + i, v)
        c.font, c.fill, c.alignment, c.border = TH_FONT, TH_FILL, CENTER, BOX


def td(ws, row, values, col=1, center=()):
    for i, v in enumerate(values):
        c = ws.cell(row, col + i, v)
        c.font, c.border = BODY, BOX
        c.alignment = CENTER if i in center else WRAP


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths):
        ws.column_dimensions[L(i + 1)].width = w


def copy_style(dst, src):
    dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
    dst.alignment, dst.number_format = copy(src.alignment), src.number_format


# ── 1. 통합매트릭스 ─────────────────────────────────────────────
# (그룹, 열 이름, 원본 열 인덱스(0-base) 또는 산출 방식)
COLS = [
    ("분류 체계 (Lv1 · Lv2 · Lv3)", "도메인 (Lv1)", 0),
    (None, "UT-ID", 1),
    (None, "위협 분류 (Lv2)", 2),
    (None, "Lv3 ID", 3),
    (None, "세부 위협 (Lv3)", 4),
    (None, "위협 설명 (Lv3)", 5),
    ("위험 평가", "발생가능성", 6),
    (None, "심각도", 7),
    (None, "위험도", "risk"),
    (None, "심각도 근거", 10),
    (None, "근거 수준", 42),
    (None, "실제 사고 수", "incidents"),
    ("교차 매핑", "주 공격 단계", 11),
    (None, "ATLAS 기법 (주)", 24),
    (None, "OWASP 매핑 (주)", "owasp"),
    ("실제 근거", "실제 사례·공격 예시", 44),
    ("통제 (Lv2)", "핵심 통제 요약", 46),
]
WIDTH = {"도메인 (Lv1)": 13, "UT-ID": 7, "위협 분류 (Lv2)": 16, "Lv3 ID": 8, "세부 위협 (Lv3)": 16,
         "위협 설명 (Lv3)": 50, "발생가능성": 7, "심각도": 7, "위험도": 8, "심각도 근거": 24, "근거 수준": 11,
         "실제 사고 수": 7, "주 공격 단계": 13, "ATLAS 기법 (주)": 22, "OWASP 매핑 (주)": 14,
         "실제 사례·공격 예시": 60, "핵심 통제 요약": 34}

sf, sv = fx["통합매트릭스"], val["통합매트릭스"]
ws = out.create_sheet("통합매트릭스")
N = sum(1 for r in sv.iter_rows(min_row=5, values_only=True) if r[3])
LAST = 4 + N
ws["A1"] = f"통합 AI 보안위협 매트릭스 v3.1 LITE — Lv1 도메인(10) · Lv2 위협 분류(36) · Lv3 세부 위협({N})"
ws["A2"] = ("기준: MITRE ATLAS v2026.09 · OWASP LLM Top 10 2026 · Agentic Top 10 2026 · "
            f"GenAI Data Security 2026 v1.0 · MCP Top 10(2025 beta) | {DATE}")
copy_style(ws["A1"], sf["A1"]); copy_style(ws["A2"], sf["A2"])

group_start = None
for j, (group, name, _) in enumerate(COLS, start=1):
    if group:
        if group_start:
            ws.merge_cells(start_row=3, start_column=group_start, end_row=3, end_column=j - 1)
        group_start = j
        ws.cell(3, j, group)
    c = ws.cell(3, j); c.font, c.fill, c.alignment, c.border = Font(bold=True, size=9, color=NAVY), GROUP_FILL, CENTER, BOX
    c = ws.cell(4, j, name); copy_style(c, sf["A4"])
    ws.column_dimensions[L(j)].width = WIDTH[name]
ws.merge_cells(start_row=3, start_column=group_start, end_row=3, end_column=len(COLS))

col_of = {name: L(j) for j, (_, name, _) in enumerate(COLS, start=1)}
G, H = col_of["발생가능성"], col_of["심각도"]
for i, r in enumerate(sv.iter_rows(min_row=5, max_row=4 + N, values_only=True)):
    row = 5 + i
    for j, (_, name, src) in enumerate(COLS, start=1):
        if src == "risk":
            v = RISK.format(g=f"{G}{row}", h=f"{H}{row}")
            src_style = sf.cell(row, 9)
        elif src == "incidents":
            v = (r[39] or 0) + (r[40] or 0)  # ATLAS 실제 사고 + OWASP 인용 실제 사고
            src_style = sf.cell(row, 40)
        elif src == "owasp":
            v = ", ".join(x for x in (r[29], r[31], r[33], r[35]) if x) or None  # LLM·ASI·DSGAI·MCP (주)
            src_style = sf.cell(row, 30)
        else:
            v = r[src]
            src_style = sf.cell(row, src + 1)
        c = ws.cell(row, j, v)
        copy_style(c, src_style)
    ws.row_dimensions[row].height = sf.row_dimensions[row].height

ws.freeze_panes = "G5"
ws.auto_filter.ref = f"A4:{L(len(COLS))}{LAST}"

# 조건부 서식·드롭다운
lvl = {"상": "C00000", "중": "7F6000", "하": "375623"}
for col in (G, H):
    for k, color in lvl.items():
        ws.conditional_formatting.add(f"{col}5:{col}{LAST}",
                                      CellIsRule(operator="equal", formula=[f'"{k}"'], font=Font(bold=True, color=color)))
R = col_of["위험도"]
for k, bg, fg in (("매우 높음", "C00000", "FFFFFF"), ("높음", "ED7D31", "FFFFFF"),
                  ("보통", "FFE699", "000000"), ("낮음", "C6E0B4", "000000")):
    ws.conditional_formatting.add(f"{R}5:{R}{LAST}", CellIsRule(
        operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", fgColor=bg), font=Font(bold=True, color=fg)))
E = col_of["근거 수준"]
for k, color in (("실제 사고 확인", "C00000"), ("실사용 기법 포함", "C55A11"),
                 ("실증·공개 취약점", "7F6000"), ("이론·시나리오", "595959")):
    ws.conditional_formatting.add(f"{E}5:{E}{LAST}",
                                  CellIsRule(operator="equal", formula=[f'"{k}"'], font=Font(bold=True, color=color)))
dv = DataValidation(type="list", formula1='"상,중,하"', allow_blank=False)
dv.add(f"{G}5:{H}{LAST}")
ws.add_data_validation(dv)

M = "통합매트릭스!"
rng = lambda name: f"{M}${col_of[name]}$5:${col_of[name]}${LAST}"

# ── 2. 개요 ─────────────────────────────────────────────────────
o = out.create_sheet("개요", 0)
o["A1"], o["A1"].font = "통합 AI 보안위협 매트릭스 (UTM) v3.1 LITE", TITLE
o["A2"], o["A2"].font = (f"Lv1 도메인 10 · Lv2 위협 분류 36 · Lv3 세부 위협 {N} · 위험 평가(발생가능성×심각도) | {DATE} "
                         "| v3.1 원본에서 핵심 시트·열만 남긴 경량 최종본"), SUB
row = 4
o.cell(row, 1, "1. 통합 대상 소스").font = SECTION; row += 1
src_o = val["개요"]
th(o, row, [c.value for c in src_o[5][:5]]); row += 1
for r in src_o.iter_rows(min_row=6, max_row=10, max_col=5, values_only=True):
    td(o, row, r); row += 1

row += 1
o.cell(row, 1, "2. 핵심 지표 (수식 연동)").font = SECTION; row += 1
th(o, row, ["지표", "값"]); row += 1
metrics = [
    ("Lv1 도메인 / Lv2 위협 분류 / Lv3 세부 위협",
     f'=SUMPRODUCT(1/COUNTIF({rng("도메인 (Lv1)")},{rng("도메인 (Lv1)")}))&" / "&'
     f'SUMPRODUCT(1/COUNTIF({rng("UT-ID")},{rng("UT-ID")}))&" / "&COUNTA({rng("Lv3 ID")})'),
    ("위험도 매우 높음 / 높음 / 보통 / 낮음",
     "=" + '&" / "&'.join(f'COUNTIF({rng("위험도")},"{k}")' for k in ("매우 높음", "높음", "보통", "낮음"))),
    ("근거 수준 실제 사고 확인 / 실사용 기법 포함 / 실증·공개 취약점 / 이론·시나리오",
     "=" + '&" / "&'.join(f'COUNTIF({rng("근거 수준")},"{k}")'
                          for k in ("실제 사고 확인", "실사용 기법 포함", "실증·공개 취약점", "이론·시나리오"))),
    ("실제 사고가 1건 이상 연결된 Lv3", f'=COUNTIF({rng("실제 사고 수")},">0")'),
]
for k, v in metrics:
    td(o, row, [k, v]); row += 1

row += 1
o.cell(row, 1, "3. 위험 매트릭스 (Lv3 수)").font = SECTION; row += 1
th(o, row, ["발생가능성 \\ 심각도", "상", "중", "하"]); row += 1
for g in ("상", "중", "하"):
    td(o, row, [g] + [f'=COUNTIFS({rng("발생가능성")},"{g}",{rng("심각도")},"{h}")' for h in ("상", "중", "하")],
       center=(0, 1, 2, 3)); row += 1

row += 1
o.cell(row, 1, "4. 시트 구성").font = SECTION; row += 1
th(o, row, ["시트", "내용"]); row += 1
for k, v in [
    ("통합매트릭스", f"핵심 산출물. Lv3 {N}행 — 분류 체계, 위험 평가, 교차 매핑(공격 단계·ATLAS·OWASP), 실제 사례, 핵심 통제(Lv2)"),
    ("도메인·평가기준", "3단 분류 정의와 작성 규칙, Lv1 도메인 리스크 정의·경계 판정 규칙, 위험 평가·근거 수준 기준"),
    ("ATLAS사례연구", "MITRE ATLAS 사례 73건 요약과 Lv3 매핑 — '실제 사례·공격 예시' 열의 CS ID 조회용"),
    ("변경이력", "버전별 변경 내역(이후 개정은 새 버전 파일로 저장하고 이 시트에 기록)"),
]:
    td(o, row, [k, v]); row += 1
widths(o, [42, 60, 30, 22, 34])

# ── 3. 도메인·평가기준 ──────────────────────────────────────────
d = out.create_sheet("도메인·평가기준")
dv_ = val["도메인·평가기준"]
d["A1"], d["A1"].font = "분류 체계 정의와 위험 평가 기준", TITLE
row = 3
d.cell(row, 1, "1. 3단 분류 체계").font = SECTION; row += 1
th(d, row, ["단계", "정의", "작성 규칙", "개수"]); row += 1
counts = [f'=SUMPRODUCT(1/COUNTIF({rng("도메인 (Lv1)")},{rng("도메인 (Lv1)")}))',
          f'=SUMPRODUCT(1/COUNTIF({rng("UT-ID")},{rng("UT-ID")}))', f'=COUNTA({rng("Lv3 ID")})']
for r, cnt in zip(dv_.iter_rows(min_row=5, max_row=7, max_col=3, values_only=True), counts):
    td(d, row, list(r) + [cnt], center=(3,)); row += 1

row += 1
d.cell(row, 1, "2. Lv1 도메인 정의").font = SECTION; row += 1
th(d, row, ["도메인", "영문명", "리스크 정의", "주요 대상 영역", "경계 판정 규칙", "Lv2 수", "Lv3 수",
            "위험도 '매우 높음'·'높음' Lv3"]); row += 1
A, B, I_ = rng("도메인 (Lv1)"), rng("UT-ID"), rng("위험도")
for r in dv_.iter_rows(min_row=11, max_row=20, max_col=5, values_only=True):
    a = f"A{row}"
    td(d, row, list(r) + [f"=SUMPRODUCT(({A}={a})/COUNTIFS({A},{A},{B},{B}))", f"=COUNTIF({A},{a})",
                          f'=COUNTIFS({A},{a},{I_},"매우 높음")+COUNTIFS({A},{a},{I_},"높음")'], center=(5, 6, 7))
    row += 1

row += 1
d.cell(row, 1, "3. 위험 평가 기준 (Lv3 단위)").font = SECTION; row += 1
th(d, row, ["항목", "상", "중", "하", "산정 방식"]); row += 1
for r in dv_.iter_rows(min_row=24, max_row=26, max_col=5, values_only=True):
    r = list(r)
    if r[0] == "발생가능성":
        r[4] = ("근거(ATLAS 사례·OWASP 인용 실제 사고·Realized 기법·공개 취약점) 기준 산정. LITE는 v3.1 산정값을 고정하며, "
                "사례·근거 개정 시 기준에 따라 재평가(드롭다운 수정)")
    if r[0] == "심각도":
        r[4] = "분석자 평가(심각도 근거 열에 사유 기재) — 기준: 통제권 범위, 정보 노출 규모, 피해 지속성. 드롭다운으로 수정 가능"
    if r[0] == "위험도":
        r[4] = "발생가능성 × 심각도 3×3 매트릭스(수식, 자동 재계산)"
    td(d, row, r); row += 1

row += 1
d.cell(row, 1, "4. 근거 수준 정의 (높은 순)").font = SECTION; row += 1
th(d, row, ["근거 수준", "판정 기준"]); row += 1
for k, v in [
    ("실제 사고 확인", "ATLAS Incident 사례 또는 OWASP 원문이 인용한 실제 사고가 1건 이상 연결('실제 사고 수' ≥ 1)"),
    ("실사용 기법 포함", "실제 사고 연결은 없으나 ATLAS Maturity 'Realized' 기법이 귀속"),
    ("실증·공개 취약점", "연구·레드팀 실증 사례(ATLAS Exercise) 또는 OWASP 인용 CVE·공개 익스플로잇 존재"),
    ("이론·시나리오", "공개 근거 없음(OWASP 위험 정의·시나리오 기반). '발생하지 않음'이 아니라 '공개 근거 없음'을 의미"),
]:
    td(d, row, [k, v]); row += 1
widths(d, [26, 30, 48, 40, 48, 8, 8, 14])

# ── 4. ATLAS사례연구 ────────────────────────────────────────────
a = out.create_sheet("ATLAS사례연구")
a["A1"], a["A1"].font = "MITRE ATLAS 사례연구(73건) — 요약 및 Lv3 매핑", TITLE
a["A2"], a["A2"].font = ("주 Lv3 = 분석자가 지정한 대표 세부 위협. 매핑 Lv3 = 공식 절차(employs) 기법의 Lv3 + 분석자 보완. "
                         "유형: Incident=실제 사고, Exercise=실증"), SUB
th(a, 4, ["사례 ID", "사례명", "유형", "발생·공개 시점", "행위자", "대상", "국문 요약", "주 Lv3", "매핑 Lv3 (전체)", "URL"])
row = 5
for r in val["ATLAS사례연구"].iter_rows(min_row=5, values_only=True):
    if not r[0]:
        continue
    td(a, row, [r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[9], r[15]], center=(2, 3, 7)); row += 1
a.freeze_panes = "B5"
a.auto_filter.ref = f"A4:J{row - 1}"
widths(a, [12, 34, 9, 9, 18, 22, 48, 8, 26, 30])

# ── 5. 변경이력 ─────────────────────────────────────────────────
h = out.create_sheet("변경이력")
h["A1"], h["A1"].font = "변경이력", TITLE
h["A2"], h["A2"].font = "개정은 새 버전 파일로 저장하고 이 시트 상단에 기록. 이전 버전 상세 이력은 v3.1 원본 '변경이력' 시트 참조", SUB
th(h, 4, ["버전", "일자", "구분", "대상", "변경 내용", "사유"])
removed_sheets = [s for s in fx.sheetnames if s not in ("개요", "통합매트릭스", "도메인·평가기준", "ATLAS사례연구", "변경이력", "설계노트")]
kept_idx = {c[2] for c in COLS if isinstance(c[2], int)} | {9 - 1, 39, 40, 29, 31, 33, 35}
removed_cols = [sf.cell(4, i + 1).value for i in range(sf.max_column) if i not in kept_idx]
log = [
    ("v3.1 LITE", DATE, "시트 삭제", f"{len(removed_sheets) + 1}개 시트",
     "삭제: " + ", ".join(removed_sheets + ["설계노트"]) + "\n(설계노트의 분류·평가 기준은 '도메인·평가기준'에 통합)",
     "최종 산출물 경량화 — 도메인·위협 분류·세부 위협·설명·사례 중심으로 재구성"),
    ("v3.1 LITE", DATE, "열 삭제", f"통합매트릭스 {sf.max_column}열 → {len(COLS)}열",
     "삭제: " + ", ".join(removed_cols),
     "패싯·세부 매핑·집계용 보조 열 제거. 원본 v3.1에 보존"),
    ("v3.1 LITE", DATE, "열 통합", "OWASP 매핑",
     "LLM·ASI·DSGAI·MCP 주 매핑 4열을 'OWASP 매핑 (주)' 1열로 통합(보조 매핑 제외)", "가독성"),
    ("v3.1 LITE", DATE, "열 통합", "실제 사고 수", "ATLAS 실제 사고 수 + OWASP 인용 실제 사고 수를 1열로 합산", "가독성"),
    ("v3.1 LITE", DATE, "수식 고정", "발생가능성·근거 수준",
     "삭제 시트를 참조하던 수식을 v3.1 산정값으로 고정. 위험도·개요 지표·도메인 집계는 LITE 내부 수식으로 재연결",
     "삭제 시트 참조 오류 방지"),
    ("v3.1", "2026-10-01", "명칭 변경", "Lv1 도메인 10개", "보호 대상 명칭 → 영역별 보안 리스크 명칭(OWASP 명명 규칙)", "—"),
    ("v3.0", "—", "구조 개편", "Lv3 127개", "Lv3 세부 위협 도입, ATLAS 기법 21개 UT 간 이동, 위험 평가 추가", "—"),
]
row = 5
for r in log:
    td(h, row, r, center=(0, 1, 2)); row += 1
widths(h, [11, 11, 10, 22, 80, 40])

for s in out.worksheets:
    s.sheet_view.showGridLines = s.title == "통합매트릭스"
out.save(OUT)
print(OUT.relative_to(ROOT), f"{N} rows", out.sheetnames, removed_sheets)
