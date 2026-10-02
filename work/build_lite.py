"""v3.1 원본에서 작업용 LITE 파일(xlsx + csv)을 생성한다. 원본은 수정하지 않는다."""
import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "reference/utm/통합_AI보안위협_매트릭스_v3.1.xlsx"
OUT_XLSX = ROOT / "work/통합_AI보안위협_매트릭스_v3.1_LITE.xlsx"
OUT_CSV = ROOT / "work/통합매트릭스_LITE.csv"

src = openpyxl.load_workbook(SRC, read_only=True, data_only=True)  # 수식은 캐시된 값으로 고정
out = openpyxl.Workbook()
out.remove(out.active)
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
WORK_FILL = PatternFill("solid", fgColor="FFF2CC")


def add_sheet(name, header, rows, widths, work_cols=0):
    ws = out.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for c in ws[1][len(header) - work_cols:]:
        c.fill, c.font = WORK_FILL, Font(bold=True)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "F2" if name == "통합매트릭스_LITE" else "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def join(*vals):
    return ", ".join(v for v in vals if v)


# 1. 통합매트릭스_LITE — 원본 열 인덱스(0-base): A~F, I, Y, AD, AF, AH, AJ, AQ, AR, AS
main_header = ["도메인 (Lv1)", "UT-ID", "위협 분류 (Lv2)", "Lv3 ID", "세부 위협 (Lv3)", "위협 설명 (Lv3)",
               "실제 사례·공격 예시", "위험도", "근거 수준", "ATLAS 기법 (주)", "OWASP 주 매핑",
               "관련 ATLAS 사례 ID", "개선 설명(안)", "개선 사례(안)", "검토 메모"]
main_rows = []
for r in src["통합매트릭스"].iter_rows(min_row=5, values_only=True):
    if not r[3]:
        continue
    main_rows.append([r[0], r[1], r[2], r[3], r[4], r[5], r[44], r[8], r[42], r[24],
                      join(r[29], r[31], r[33], r[35]), r[43], None, None, None])
add_sheet("통합매트릭스_LITE", main_header, main_rows,
          [22, 8, 22, 9, 22, 55, 60, 9, 13, 25, 18, 25, 55, 60, 30], work_cols=3)

# 2. 도메인 정의 — 도메인·평가기준 2번 표(수식 열 제외)
dom_rows = [list(r[:5]) for r in src["도메인·평가기준"].iter_rows(min_row=11, max_row=20, values_only=True)]
add_sheet("도메인 정의", ["도메인", "영문명", "리스크 정의", "주요 대상 영역", "경계 판정 규칙"],
          dom_rows, [28, 30, 55, 40, 45])

# 3. ATLAS 사례 — 사례 교체·검증용 조회표
cs_rows = [[r[0], r[1], r[2], r[3], r[4], r[6], r[7], r[9], r[14]]
           for r in src["ATLAS사례연구"].iter_rows(min_row=5, values_only=True) if r[0]]
add_sheet("ATLAS 사례", ["사례 ID", "사례명", "유형", "발생·공개 시점", "행위자", "국문 요약",
                        "주 Lv3", "매핑 Lv3 (전체)", "분석↔공식 정합성"],
          cs_rows, [12, 35, 10, 10, 20, 50, 9, 30, 14])

# 4. OWASP 원문 — 시나리오·CVE 사례 근거용
ow_rows = [[r[0], r[1], r[3], r[4], r[6], r[7]]
           for r in src["OWASP상세"].iter_rows(min_row=5, values_only=True) if r[0]]
add_sheet("OWASP 원문", ["소스", "코드", "유형", "번호", "원문 내용", "주 매핑 UT"],
          ow_rows, [16, 9, 12, 6, 90, 14])

# 5. 작성 기준
rules = [
    ["원본", "reference/utm/통합_AI보안위협_매트릭스_v3.1.xlsx (수식은 캐시 값으로 고정, 원본 미수정)"],
    ["Lv1 도메인", "OWASP 위협·리스크 명명 규칙: [영역] + [위협 명사], 대등한 두 위협만 '및'으로 결합"],
    ["Lv2 위협 분류", "'·' 나열 없이 하나의 개념으로 명명. UT-ID 불변"],
    ["Lv3 세부 위협", "① 독립 메커니즘 ② 근거 1개 이상(ATLAS 기법·OWASP 예시·사례) ③ 공격 단계·생애주기·통제 중 하나 이상 구별"],
    ["위협 설명(v3.1)", "'수단→행위→결과' 명사형 종결"],
    ["사례 표기(v3.1)", "• [ATLAS CSxxxx·실제사고|실증] … / • [OWASP 인용] … / • [OWASP 시나리오] …"],
    ["작업 열", "노란 머리글(개선 설명(안)·개선 사례(안)·검토 메모)에 개선안 작성 → 확정 후 원본 v3.x에 반영"],
]
add_sheet("작성 기준", ["항목", "내용"], rules, [18, 110])

out.save(OUT_XLSX)

# 작업 중 읽기용 CSV(사례 ID 목록 등 긴 참조 열 제외)
with OUT_CSV.open("w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(main_header[:11])
    for r in main_rows:
        w.writerow(r[:11])
print(f"{len(main_rows)} Lv3 rows, {len(cs_rows)} cases, {len(ow_rows)} OWASP items")
